import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import sqlite3
from pathlib import Path
from datetime import datetime

APP_NAME = "3D Print Cost"

def get_app_data_dir():
    """Cartella persistente per impostazioni e database."""
    import os
    if os.name == "nt":
        base = Path(os.environ.get("LOCALAPPDATA", Path.home()))
        folder = base / "3D Print Cost"
    else:
        folder = Path.home() / ".3d_print_cost"
    folder.mkdir(parents=True, exist_ok=True)
    return folder

APP_DATA_DIR = get_app_data_dir()
DB_PATH = APP_DATA_DIR / "costi_stampa_3d.db"

DEFAULT_ELECTRICITY = 0.15
DEFAULT_POWER_W = 200.0


class Database:
    def __init__(self, path):
        self.conn = sqlite3.connect(path)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
        """)
        self.conn.execute("""
            CREATE TABLE IF NOT EXISTS filaments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT UNIQUE NOT NULL COLLATE NOCASE,
                price_per_kg REAL NOT NULL CHECK(price_per_kg >= 0)
            )
        """)
        self.conn.commit()

        if self.get_setting("electricity_price") is None:
            self.set_setting("electricity_price", DEFAULT_ELECTRICITY)
        if self.get_setting("printer_power_w") is None:
            self.set_setting("printer_power_w", DEFAULT_POWER_W)

    def get_setting(self, key):
        row = self.conn.execute(
            "SELECT value FROM settings WHERE key = ?", (key,)
        ).fetchone()
        return row[0] if row else None

    def set_setting(self, key, value):
        self.conn.execute(
            """INSERT INTO settings(key, value) VALUES(?, ?)
               ON CONFLICT(key) DO UPDATE SET value=excluded.value""",
            (key, str(value))
        )
        self.conn.commit()

    def filaments(self):
        return self.conn.execute(
            "SELECT id, name, price_per_kg FROM filaments ORDER BY name COLLATE NOCASE"
        ).fetchall()

    def add_filament(self, name, price):
        self.conn.execute(
            "INSERT INTO filaments(name, price_per_kg) VALUES(?, ?)",
            (name.strip(), price)
        )
        self.conn.commit()

    def update_filament(self, filament_id, name, price):
        self.conn.execute(
            "UPDATE filaments SET name=?, price_per_kg=? WHERE id=?",
            (name.strip(), price, filament_id)
        )
        self.conn.commit()

    def delete_filament(self, filament_id):
        self.conn.execute("DELETE FROM filaments WHERE id=?", (filament_id,))
        self.conn.commit()


def parse_number(value, field, invalid_template='Il campo "{field}" deve contenere un numero valido.'):
    try:
        return float(value.strip().replace(",", "."))
    except ValueError:
        raise ValueError(invalid_template.format(field=field))


def euro(value):
    return f"€ {value:.2f}".replace(".", ",")



TRANSLATIONS = {
"it": {
"subtitle":"Calcola e archivia il costo reale dei tuoi piatti di stampa","print_data":"DATI DELLA STAMPA",
"filament":"Filamento","total_filament":"Filamento totale (g)","hours":"Ore","minutes":"Minuti","calculate":"CALCOLA COSTO",
"summary":"RIEPILOGO","material_blank":"Costo filamento: —","energy_cost_blank":"Costo corrente: —","kwh_blank":"Energia stimata: —",
"total_cost":"COSTO TOTALE","save_excel":"SALVA RISULTATO IN EXCEL","manage_filaments":"Gestisci filamenti",
"current_settings":"Impostazioni corrente","status":"Corrente: {elec:.3f} €/kWh   •   Consumo medio: {power:.0f} W   •   Database locale: {db}",
"first_run":"Primo avvio","first_run_msg":"Il database dei filamenti è vuoto.\n\nClicca su “Gestisci filamenti” e aggiungi, ad esempio:\nAnycubic PLA — 20 €/kg.\n\nIl prezzo della corrente è già impostato a 0,15 €/kWh.",
"filament_manager":"Gestione filamenti","price_kg":"Prezzo €/kg","name":"Nome:","add":"Aggiungi","save_edit":"Salva modifica",
"delete":"Elimina","new_clear":"Nuovo / Pulisci","error":"Errore","price":"Prezzo","enter_filament_name":"Inserisci il nome del filamento.",
"price_negative":"Il prezzo non può essere negativo.","filament_exists":"Esiste già un filamento con questo nome.","selection":"Selezione",
"select_edit":"Seleziona prima un filamento da modificare.","select_delete":"Seleziona il filamento da eliminare.","confirm":"Conferma",
"delete_confirm":"Eliminare “{name}” dal database?","settings":"Impostazioni","electricity_cost":"Costo corrente (€/kWh):",
"avg_power":"Consumo medio stampante (W):","power_help":"Il consumo medio serve per convertire le ore di stampa in kWh.\nValore iniziale impostato: 200 W.",
"save_settings":"Salva impostazioni","electricity_field":"Costo corrente","power_field":"Consumo medio","values_negative":"I valori non possono essere negativi.",
"saved":"Salvato","settings_saved":"Impostazioni salvate.","invalid_number":"Il campo “{field}” deve contenere un numero valido.",
"no_result":"Nessun risultato","no_result_msg":"Calcola prima il costo del piatto, poi premi “Salva risultato in Excel”.",
"save_result":"Salva risultato","excel_question":"Vuoi aggiungere questo piatto a un file Excel GIÀ ESISTENTE?\n\nSì = scegli un file esistente\nNo = crea un nuovo file Excel\nAnnulla = non salvare",
"select_excel":"Seleziona il file Excel da aggiornare","excel_files":"File Excel","plate_added":"Piatto aggiunto correttamente a:\n{path}",
"excel_error":"Errore Excel","excel_update_error":"Non è stato possibile aggiornare il file Excel.\n\nAssicurati che il file non sia aperto in Excel e che sia un file .xlsx valido.\n\nDettagli: {error}",
"create_excel":"Crea un nuovo riepilogo Excel","summary_created":"Nuovo riepilogo creato correttamente:\n{path}",
"excel_create_error":"Non è stato possibile creare il file Excel.\n\nDettagli: {error}",
"excel_headers":["Piatto","Data/Ora","Filamento","Filamento totale (g)","Prezzo filamento (€/kg)","Ore stampa","Consumo medio (W)","Energia stimata (kWh)","Prezzo corrente (€/kWh)","Costo filamento (€)","Costo corrente (€)","Costo totale piatto (€)"],
"excel_sheet":"Riepilogo stampa","excel_title":"RIEPILOGO COSTI STAMPA 3D","object_total":"TOTALE OGGETTO:",
"openpyxl_missing":"Per l'esportazione Excel serve il pacchetto 'openpyxl'. Installa una sola volta con: pip install openpyxl",
"no_filament":"Nessun filamento selezionato. Aggiungine uno in Gestisci filamenti.","total_filament_field":"Filamento totale",
"negative_inputs":"Grammi, ore e minuti non possono essere negativi.","minutes_range":"I minuti devono essere compresi tra 0 e 59.",
"filament_not_found":"Filamento non trovato nel database.","material_result":"Costo filamento: {cost}  ({grams:.2f} g × {price:.2f} €/kg)",
"kwh_result":"Energia stimata: {kwh:.3f} kWh  ({power:.0f} W × {hours:.2f} h)",
"energy_result":"Costo corrente: {cost}  ({electricity:.3f} €/kWh)"
},
"en": {
"subtitle":"Calculate and save the real cost of your 3D printing plates","print_data":"PRINT DATA",
"filament":"Filament","total_filament":"Total filament (g)","hours":"Hours","minutes":"Minutes","calculate":"CALCULATE COST",
"summary":"SUMMARY","material_blank":"Filament cost: —","energy_cost_blank":"Electricity cost: —","kwh_blank":"Estimated energy: —",
"total_cost":"TOTAL COST","save_excel":"SAVE RESULT TO EXCEL","manage_filaments":"Manage filaments",
"current_settings":"Electricity settings","status":"Electricity: €{elec:.3f}/kWh   •   Average power: {power:.0f} W   •   Local database: {db}",
"first_run":"First run","first_run_msg":"The filament database is empty.\n\nClick “Manage filaments” and add, for example:\nAnycubic PLA — €20/kg.\n\nElectricity is already set to €0.15/kWh.",
"filament_manager":"Filament manager","price_kg":"Price €/kg","name":"Name:","add":"Add","save_edit":"Save changes",
"delete":"Delete","new_clear":"New / Clear","error":"Error","price":"Price","enter_filament_name":"Enter the filament name.",
"price_negative":"The price cannot be negative.","filament_exists":"A filament with this name already exists.","selection":"Selection",
"select_edit":"Select a filament to edit first.","select_delete":"Select the filament to delete.","confirm":"Confirm",
"delete_confirm":"Delete “{name}” from the database?","settings":"Settings","electricity_cost":"Electricity price (€/kWh):",
"avg_power":"Average printer consumption (W):","power_help":"Average power is used to convert print time into kWh.\nInitial default value: 200 W.",
"save_settings":"Save settings","electricity_field":"Electricity price","power_field":"Average power","values_negative":"Values cannot be negative.",
"saved":"Saved","settings_saved":"Settings saved.","invalid_number":"The “{field}” field must contain a valid number.",
"no_result":"No result","no_result_msg":"Calculate the plate cost first, then click “Save result to Excel”.",
"save_result":"Save result","excel_question":"Do you want to add this plate to an EXISTING Excel file?\n\nYes = select an existing file\nNo = create a new Excel file\nCancel = do not save",
"select_excel":"Select the Excel file to update","excel_files":"Excel files","plate_added":"Plate successfully added to:\n{path}",
"excel_error":"Excel error","excel_update_error":"The Excel file could not be updated.\n\nMake sure the file is not open in Excel and is a valid .xlsx file.\n\nDetails: {error}",
"create_excel":"Create a new Excel summary","summary_created":"New summary successfully created:\n{path}",
"excel_create_error":"The Excel file could not be created.\n\nDetails: {error}",
"excel_headers":["Plate","Date/Time","Filament","Total filament (g)","Filament price (€/kg)","Print hours","Average power (W)","Estimated energy (kWh)","Electricity price (€/kWh)","Filament cost (€)","Electricity cost (€)","Total plate cost (€)"],
"excel_sheet":"Print summary","excel_title":"3D PRINTING COST SUMMARY","object_total":"OBJECT TOTAL:",
"openpyxl_missing":"Excel export requires the 'openpyxl' package. Install it once with: pip install openpyxl",
"no_filament":"No filament selected. Add one in Manage filaments.","total_filament_field":"Total filament",
"negative_inputs":"Grams, hours and minutes cannot be negative.","minutes_range":"Minutes must be between 0 and 59.",
"filament_not_found":"Filament not found in the database.","material_result":"Filament cost: {cost}  ({grams:.2f} g × €{price:.2f}/kg)",
"kwh_result":"Estimated energy: {kwh:.3f} kWh  ({power:.0f} W × {hours:.2f} h)",
"energy_result":"Electricity cost: {cost}  (€{electricity:.3f}/kWh)"
}
}

class FilamentManager(tk.Toplevel):
    def __init__(self, parent, db, on_change, lang="it"):
        super().__init__(parent)
        self.db, self.on_change, self.lang = db, on_change, lang
        T = TRANSLATIONS[lang]
        self.title(T["filament_manager"])
        self.configure(bg="#111827")
        self.geometry("780x540")
        self.minsize(680, 470)
        self.transient(parent)

        frame = ttk.Frame(self, padding=22, style="Card.TFrame")
        frame.pack(fill="both", expand=True)

        self.tree = ttk.Treeview(frame, columns=("name", "price"), show="headings", height=13)
        self.tree.heading("name", text=T["filament"])
        self.tree.heading("price", text=T["price_kg"])
        self.tree.column("name", width=420)
        self.tree.column("price", width=150, anchor="center")
        self.tree.pack(fill="both", expand=True)
        self.tree.bind("<<TreeviewSelect>>", self.load_selected)

        form = ttk.Frame(frame, style="Card.TFrame")
        form.pack(fill="x", pady=(16, 8))
        ttk.Label(form, text=T["name"], style="Card.TLabel").grid(row=0, column=0, sticky="w")
        self.name_var = tk.StringVar()
        ttk.Entry(form, textvariable=self.name_var, width=38, style="Modern.TEntry").grid(
            row=0, column=1, padx=(8, 18), sticky="ew")
        ttk.Label(form, text="€/kg:", style="Card.TLabel").grid(row=0, column=2, sticky="w")
        self.price_var = tk.StringVar()
        ttk.Entry(form, textvariable=self.price_var, width=14, style="Modern.TEntry").grid(
            row=0, column=3, padx=(8, 0))
        form.columnconfigure(1, weight=1)

        buttons = ttk.Frame(frame, style="Card.TFrame")
        buttons.pack(fill="x", pady=(8, 0))
        ttk.Button(buttons, text=T["add"], command=self.add, style="Secondary.TButton").pack(side="left")
        ttk.Button(buttons, text=T["save_edit"], command=self.update, style="Secondary.TButton").pack(side="left", padx=7)
        ttk.Button(buttons, text=T["delete"], command=self.delete, style="Secondary.TButton").pack(side="left")
        ttk.Button(buttons, text=T["new_clear"], command=self.clear, style="Secondary.TButton").pack(side="right")
        self.refresh()

    def refresh(self):
        for item in self.tree.get_children(): self.tree.delete(item)
        for fid, name, price in self.db.filaments():
            self.tree.insert("", "end", iid=str(fid), values=(name, f"{price:.2f}"))
        self.on_change()

    def load_selected(self, _event=None):
        selected = self.tree.selection()
        if selected:
            values = self.tree.item(selected[0], "values")
            self.name_var.set(values[0]); self.price_var.set(str(values[1]).replace(".", ","))

    def clear(self):
        self.tree.selection_remove(self.tree.selection()); self.name_var.set(""); self.price_var.set("")

    def add(self):
        T = TRANSLATIONS[self.lang]
        name = self.name_var.get().strip()
        if not name:
            messagebox.showerror(T["error"], T["enter_filament_name"], parent=self); return
        try:
            price = parse_number(self.price_var.get(), T["price"], T["invalid_number"])
            if price < 0: raise ValueError(T["price_negative"])
            self.db.add_filament(name, price); self.clear(); self.refresh()
        except sqlite3.IntegrityError:
            messagebox.showerror(T["error"], T["filament_exists"], parent=self)
        except ValueError as e:
            messagebox.showerror(T["error"], str(e), parent=self)

    def update(self):
        T = TRANSLATIONS[self.lang]
        selected = self.tree.selection()
        if not selected:
            messagebox.showinfo(T["selection"], T["select_edit"], parent=self); return
        try:
            name = self.name_var.get().strip()
            if not name: raise ValueError(T["enter_filament_name"])
            price = parse_number(self.price_var.get(), T["price"], T["invalid_number"])
            if price < 0: raise ValueError(T["price_negative"])
            self.db.update_filament(int(selected[0]), name, price); self.refresh()
        except sqlite3.IntegrityError:
            messagebox.showerror(T["error"], T["filament_exists"], parent=self)
        except ValueError as e:
            messagebox.showerror(T["error"], str(e), parent=self)

    def delete(self):
        T = TRANSLATIONS[self.lang]
        selected = self.tree.selection()
        if not selected:
            messagebox.showinfo(T["selection"], T["select_delete"], parent=self); return
        name = self.tree.item(selected[0], "values")[0]
        if messagebox.askyesno(T["confirm"], T["delete_confirm"].format(name=name), parent=self):
            self.db.delete_filament(int(selected[0])); self.clear(); self.refresh()


class SettingsWindow(tk.Toplevel):
    def __init__(self, parent, db, on_change, lang="it"):
        super().__init__(parent)
        self.db, self.on_change, self.lang = db, on_change, lang
        T = TRANSLATIONS[lang]
        self.title(T["settings"])
        self.configure(bg="#111827")
        self.geometry("570x320")
        self.resizable(False, False)
        self.transient(parent)

        f = ttk.Frame(self, padding=24, style="Card.TFrame"); f.pack(fill="both", expand=True)
        ttk.Label(f, text=T["electricity_cost"], style="Card.TLabel").grid(row=0, column=0, sticky="w", pady=10)
        self.elec = tk.StringVar(value=str(db.get_setting("electricity_price")).replace(".", ","))
        ttk.Entry(f, textvariable=self.elec, width=16, style="Modern.TEntry").grid(row=0, column=1, pady=10)

        ttk.Label(f, text=T["avg_power"], style="Card.TLabel").grid(row=1, column=0, sticky="w", pady=10)
        self.power = tk.StringVar(value=str(db.get_setting("printer_power_w")).replace(".", ","))
        ttk.Entry(f, textvariable=self.power, width=16, style="Modern.TEntry").grid(row=1, column=1, pady=10)

        ttk.Label(f, text=T["power_help"], style="Muted.TLabel").grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(10, 18))
        ttk.Button(f, text=T["save_settings"], command=self.save, style="Primary.TButton").grid(
            row=3, column=0, columnspan=2, sticky="ew")

    def save(self):
        T = TRANSLATIONS[self.lang]
        try:
            elec = parse_number(self.elec.get(), T["electricity_field"], T["invalid_number"])
            power = parse_number(self.power.get(), T["power_field"], T["invalid_number"])
            if elec < 0 or power < 0: raise ValueError(T["values_negative"])
            self.db.set_setting("electricity_price", elec); self.db.set_setting("printer_power_w", power)
            self.on_change()
            messagebox.showinfo(T["saved"], T["settings_saved"], parent=self); self.destroy()
        except ValueError as e:
            messagebox.showerror(T["error"], str(e), parent=self)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.db = Database(DB_PATH)
        self.withdraw()
        self.lang = self.choose_language()
        self.db.set_setting("language", self.lang)
        self.T = TRANSLATIONS[self.lang]

        self.title(APP_NAME)
        self.geometry("1120x820")
        self.minsize(980, 740)
        self.configure(bg="#0B1120")

        # Dopo la scelta della lingua, apre l'app massimizzata sul desktop.
        try:
            self.state("zoomed")
        except tk.TclError:
            sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
            self.geometry(f"{sw}x{sh}+0+0")

        try:
            import sys
            base_dir = Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).parent))
            ico = base_dir / "3D_Print_Cost.ico"
            if ico.exists(): self.iconbitmap(str(ico))
        except Exception:
            pass

        self.setup_styles()
        self.build_ui()
        self.refresh()
        self.deiconify()

        if not self.db.filaments():
            self.after(300, lambda: messagebox.showinfo(
                self.T["first_run"], self.T["first_run_msg"], parent=self))

    def choose_language(self):
        dlg = tk.Toplevel(self)
        dlg.title("3D Print Cost — Language / Lingua")
        dlg.configure(bg="#0B1120")
        dlg.geometry("520x330")
        dlg.resizable(False, False)
        dlg.grab_set()
        dlg.protocol("WM_DELETE_WINDOW", lambda: choose("it"))
        result = {"lang": None}

        def choose(lang):
            result["lang"] = lang
            dlg.destroy()

        tk.Label(dlg, text="3D PRINT COST", bg="#0B1120", fg="#F8FAFC",
                 font=("Bahnschrift SemiBold", 27)).pack(pady=(35, 8))
        tk.Label(dlg, text="Scegli la lingua  •  Choose your language",
                 bg="#0B1120", fg="#94A3B8",
                 font=("Bahnschrift", 12)).pack(pady=(0, 28))

        btn_frame = tk.Frame(dlg, bg="#0B1120"); btn_frame.pack()
        for text, lang in [("🇮🇹  Italiano", "it"), ("🇬🇧  English", "en")]:
            tk.Button(btn_frame, text=text, command=lambda x=lang: choose(x),
                      font=("Bahnschrift SemiBold", 12),
                      bg="#1D2A3D", fg="#F8FAFC", activebackground="#293A52",
                      activeforeground="#FFFFFF", relief="flat", cursor="hand2",
                      width=17, height=2).pack(side="left", padx=10)
        tk.Label(dlg, text="La scelta può essere cambiata a ogni avvio.\nYou can choose again every time the app starts.",
                 bg="#0B1120", fg="#64748B", font=("Bahnschrift", 10)).pack(pady=(28, 0))
        self.wait_window(dlg)
        return result["lang"] or "it"

    def setup_styles(self):
        style = ttk.Style(self)
        try: style.theme_use("clam")
        except tk.TclError: pass
        BG, CARD, CARD2, TEXT, MUTED = "#0B1120", "#151E2E", "#101827", "#F8FAFC", "#94A3B8"
        ACCENT, ACCENT_HOVER, BLUE, BORDER, ENTRY = "#2DD4BF", "#14B8A6", "#38BDF8", "#263449", "#0A1220"

        # Bahnschrift gives the interface a modern technical look and is included with current Windows versions.
        UI = "Bahnschrift"
        DISPLAY = "Bahnschrift SemiBold"
        style.configure(".", font=(UI, 12))
        style.configure("App.TFrame", background=BG)
        style.configure("Card.TFrame", background=CARD, relief="flat")
        style.configure("Card2.TFrame", background=CARD2, relief="flat")
        style.configure("Title.TLabel", background=BG, foreground=TEXT, font=(DISPLAY, 32))
        style.configure("Subtitle.TLabel", background=BG, foreground=MUTED, font=(UI, 12))
        style.configure("CardTitle.TLabel", background=CARD, foreground=TEXT, font=(DISPLAY, 15))
        style.configure("Card.TLabel", background=CARD, foreground=TEXT, font=(UI, 11))
        style.configure("Muted.TLabel", background=CARD, foreground=MUTED, font=(UI, 10))
        style.configure("Status.TLabel", background=BG, foreground=MUTED, font=(UI, 10))
        style.configure("TotalCaption.TLabel", background=CARD2, foreground=MUTED, font=(UI, 11, "bold"))
        style.configure("Total.TLabel", background=CARD2, foreground="#5EEAD4", font=(DISPLAY, 36))
        style.configure("Modern.TEntry", fieldbackground=ENTRY, foreground=TEXT, bordercolor=BORDER,
                        lightcolor=BORDER, darkcolor=BORDER, padding=10, font=(UI, 11))
        style.configure("Modern.TCombobox", fieldbackground=ENTRY, background=ENTRY, foreground=TEXT,
                        arrowcolor=TEXT, bordercolor=BORDER, lightcolor=BORDER, darkcolor=BORDER,
                        padding=9, font=(UI, 11))
        style.map("Modern.TCombobox", fieldbackground=[("readonly", ENTRY)], foreground=[("readonly", TEXT)],
                  selectbackground=[("readonly", ENTRY)], selectforeground=[("readonly", TEXT)])
        style.configure("Primary.TButton", background=ACCENT, foreground="#07140B",
                        font=(DISPLAY, 12), padding=(16, 13), borderwidth=0)
        style.map("Primary.TButton", background=[("active", ACCENT_HOVER), ("pressed", ACCENT_HOVER)])
        style.configure("Secondary.TButton", background="#1D2A3D", foreground=TEXT,
                        font=(DISPLAY, 11), padding=(14, 11), borderwidth=0)
        style.map("Secondary.TButton", background=[("active", "#293A52"), ("pressed", "#293A52")])
        style.configure("Excel.TButton", background=BLUE, foreground="white",
                        font=(DISPLAY, 11), padding=(14, 11), borderwidth=0)
        style.map("Excel.TButton", background=[("active", "#0EA5E9"), ("pressed", "#0EA5E9")])
        style.configure("Treeview", background=ENTRY, fieldbackground=ENTRY, foreground=TEXT,
                        rowheight=30, font=(UI, 10))
        style.configure("Treeview.Heading", background="#1D2A3D", foreground=TEXT, font=(DISPLAY, 10))
        self.option_add("*TCombobox*Listbox.background", ENTRY)
        self.option_add("*TCombobox*Listbox.foreground", TEXT)
        self.option_add("*TCombobox*Listbox.selectBackground", "#263449")
        self.option_add("*TCombobox*Listbox.selectForeground", TEXT)

    def build_ui(self):
        T = self.T
        outer = ttk.Frame(self, padding=(54, 38), style="App.TFrame"); outer.pack(fill="both", expand=True)
        header = ttk.Frame(outer, style="App.TFrame"); header.pack(fill="x", pady=(0, 30))
        ttk.Label(header, text="3D PRINT COST", style="Title.TLabel").pack(anchor="w")
        ttk.Label(header, text=T["subtitle"], style="Subtitle.TLabel").pack(anchor="w", pady=(4, 0))
        accent_bar = tk.Frame(header, bg="#2DD4BF", height=3, bd=0, highlightthickness=0)
        accent_bar.pack(fill="x", pady=(18, 0))

        body = ttk.Frame(outer, style="App.TFrame"); body.pack(fill="both", expand=True)
        body.columnconfigure(0, weight=3); body.columnconfigure(1, weight=2); body.rowconfigure(0, weight=1)

        input_card = ttk.Frame(body, padding=32, style="Card.TFrame")
        input_card.grid(row=0, column=0, sticky="nsew", padx=(0, 16)); input_card.columnconfigure((0,1), weight=1)
        ttk.Label(input_card, text=T["print_data"], style="CardTitle.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 22))
        ttk.Label(input_card, text=T["filament"], style="Muted.TLabel").grid(row=1, column=0, columnspan=2, sticky="w")
        self.filament_var = tk.StringVar()
        self.filament_combo = ttk.Combobox(input_card, textvariable=self.filament_var, state="readonly", style="Modern.TCombobox")
        self.filament_combo.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(7, 18))
        ttk.Label(input_card, text=T["total_filament"], style="Muted.TLabel").grid(row=3, column=0, columnspan=2, sticky="w")
        self.grams = tk.StringVar()
        ttk.Entry(input_card, textvariable=self.grams, style="Modern.TEntry").grid(row=4, column=0, columnspan=2, sticky="ew", pady=(7, 18))
        ttk.Label(input_card, text=T["hours"], style="Muted.TLabel").grid(row=5, column=0, sticky="w", padx=(0, 9))
        ttk.Label(input_card, text=T["minutes"], style="Muted.TLabel").grid(row=5, column=1, sticky="w", padx=(9, 0))
        self.hours, self.minutes = tk.StringVar(value="0"), tk.StringVar(value="0")
        ttk.Entry(input_card, textvariable=self.hours, style="Modern.TEntry").grid(row=6, column=0, sticky="ew", padx=(0,9), pady=(7,24))
        ttk.Entry(input_card, textvariable=self.minutes, style="Modern.TEntry").grid(row=6, column=1, sticky="ew", padx=(9,0), pady=(7,24))
        ttk.Button(input_card, text="◆  " + T["calculate"], command=self.calculate, style="Primary.TButton").grid(row=7, column=0, columnspan=2, sticky="ew")

        result_card = ttk.Frame(body, padding=32, style="Card.TFrame")
        result_card.grid(row=0, column=1, sticky="nsew", padx=(16,0))
        ttk.Label(result_card, text=T["summary"], style="CardTitle.TLabel").pack(anchor="w", pady=(0,20))
        self.material_result = tk.StringVar(value=T["material_blank"])
        self.energy_result = tk.StringVar(value=T["energy_cost_blank"])
        self.kwh_result = tk.StringVar(value=T["kwh_blank"])
        self.total_result = tk.StringVar(value="—")
        for var in (self.material_result, self.kwh_result, self.energy_result):
            ttk.Label(result_card, textvariable=var, style="Card.TLabel", wraplength=380).pack(anchor="w", pady=7)
        total_card = ttk.Frame(result_card, padding=24, style="Card2.TFrame"); total_card.pack(fill="x", pady=(25,18))
        ttk.Label(total_card, text=T["total_cost"], style="TotalCaption.TLabel").pack(anchor="w")
        ttk.Label(total_card, textvariable=self.total_result, style="Total.TLabel").pack(anchor="w", pady=(4,0))
        ttk.Button(result_card, text="▣  " + T["save_excel"], command=self.save_result_excel, style="Excel.TButton").pack(fill="x", pady=(8,0))

        actions = ttk.Frame(outer, padding=(0,22,0,0), style="App.TFrame"); actions.pack(fill="x")
        ttk.Button(actions, text="◉  " + T["manage_filaments"], command=lambda: FilamentManager(self, self.db, self.refresh, self.lang), style="Secondary.TButton").pack(side="left")
        ttk.Button(actions, text="⚙  " + T["current_settings"], command=lambda: SettingsWindow(self, self.db, self.refresh, self.lang), style="Secondary.TButton").pack(side="left", padx=10)
        self.status = tk.StringVar()
        ttk.Label(outer, textvariable=self.status, style="Status.TLabel").pack(anchor="w", pady=(16,0))

    def refresh(self):
        rows = self.db.filaments(); names = [r[1] for r in rows]; current = self.filament_var.get()
        self.filament_combo["values"] = names
        self.filament_var.set(current if current in names else (names[0] if names else ""))
        elec = float(self.db.get_setting("electricity_price")); power = float(self.db.get_setting("printer_power_w"))
        self.status.set(self.T["status"].format(elec=elec, power=power, db=DB_PATH.name))

    def _last_result(self): return getattr(self, "last_result", None)

    def save_result_excel(self):
        T = self.T; result = self._last_result()
        if not result:
            messagebox.showinfo(T["no_result"], T["no_result_msg"], parent=self); return
        answer = messagebox.askyesnocancel(T["save_result"], T["excel_question"], parent=self)
        if answer is None: return
        if answer:
            path = filedialog.askopenfilename(parent=self, title=T["select_excel"], filetypes=[(T["excel_files"], "*.xlsx")])
            if not path: return
            try:
                self._append_to_excel(Path(path), result)
                messagebox.showinfo(T["saved"], T["plate_added"].format(path=path), parent=self)
            except Exception as e:
                messagebox.showerror(T["excel_error"], T["excel_update_error"].format(error=e), parent=self)
        else:
            path = filedialog.asksaveasfilename(parent=self, title=T["create_excel"], defaultextension=".xlsx",
                                                filetypes=[(T["excel_files"], "*.xlsx")], initialfile="3d_print_cost_summary.xlsx")
            if not path: return
            try:
                self._create_excel(Path(path), result)
                messagebox.showinfo(T["saved"], T["summary_created"].format(path=path), parent=self)
            except Exception as e:
                messagebox.showerror(T["excel_error"], T["excel_create_error"].format(error=e), parent=self)

    def _excel_headers(self):
        return TRANSLATIONS[self.lang]["excel_headers"]

    def _result_row(self, result, plate_number):
        return [plate_number, datetime.now().strftime("%d/%m/%Y %H:%M"), result["filament"], result["grams"],
                result["price_per_kg"], result["total_hours"], result["power_w"], result["kwh"], result["electricity"],
                result["material_cost"], result["energy_cost"], result["total"]]

    def _create_excel(self, path, result):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise RuntimeError(self.T["openpyxl_missing"])
        wb = Workbook(); ws = wb.active; ws.title = self.T["excel_sheet"]
        ws.merge_cells("A1:L1"); ws["A1"] = self.T["excel_title"]; ws["A1"].font = Font(size=16, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center")
        for col, value in enumerate(self._excel_headers(), 1):
            cell = ws.cell(row=3, column=col, value=value); cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7"); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col, value in enumerate(self._result_row(result, 1), 1): ws.cell(row=4, column=col, value=value)
        ws["K6"] = self.T["object_total"]; ws["K6"].font = Font(bold=True); ws["L6"] = "=SUM(L4:L4)"; ws["L6"].font = Font(size=14, bold=True)
        self._format_excel_sheet(ws); wb.save(path)

    def _append_to_excel(self, path, result):
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Font
        except ImportError:
            raise RuntimeError(self.T["openpyxl_missing"])
        wb = load_workbook(path)
        # Compatible with files previously generated in either language.
        candidates = ["Riepilogo stampa", "Print summary", self.T["excel_sheet"]]
        ws = next((wb[n] for n in candidates if n in wb.sheetnames), wb.active)
        row_num = 4
        while ws.cell(row=row_num, column=1).value not in (None, ""): row_num += 1
        for col, value in enumerate(self._result_row(result, row_num - 3), 1): ws.cell(row=row_num, column=col, value=value)
        total_labels = {"TOTALE OGGETTO:", "OBJECT TOTAL:"}
        for r in range(4, ws.max_row + 3):
            if ws.cell(r,11).value in total_labels:
                ws.cell(r,11).value = None; ws.cell(r,12).value = None
        total_row = row_num + 2
        ws.cell(total_row,11,self.T["object_total"]); ws.cell(total_row,11).font = Font(bold=True)
        ws.cell(total_row,12,f"=SUM(L4:L{row_num})"); ws.cell(total_row,12).font = Font(size=14,bold=True)
        self._format_excel_sheet(ws); wb.save(path)

    def _format_excel_sheet(self, ws):
        try:
            from openpyxl.styles import Alignment, Border, Side
        except ImportError: return
        thin = Side(style="thin", color="D9D9D9")
        for row in range(4, ws.max_row + 1):
            if isinstance(ws.cell(row,1).value,(int,float)):
                ws.cell(row,4).number_format='0.00'; ws.cell(row,5).number_format='€ 0.00'
                ws.cell(row,6).number_format='0.00'; ws.cell(row,7).number_format='0'
                ws.cell(row,8).number_format='0.000'; ws.cell(row,9).number_format='€ 0.000'
                for col in (10,11,12): ws.cell(row,col).number_format='€ 0.00'
                for col in range(1,13):
                    ws.cell(row,col).border=Border(bottom=thin); ws.cell(row,col).alignment=Alignment(vertical="center")
        for row in range(4, ws.max_row + 1):
            if ws.cell(row,11).value in {"TOTALE OGGETTO:", "OBJECT TOTAL:"}: ws.cell(row,12).number_format='€ 0.00'
        widths={"A":10,"B":19,"C":28,"D":20,"E":22,"F":14,"G":19,"H":22,"I":22,"J":20,"K":19,"L":22}
        for col,width in widths.items(): ws.column_dimensions[col].width=width
        ws.freeze_panes="A4"; ws.auto_filter.ref=f"A3:L{max(4,ws.max_row)}"

    def calculate(self):
        T = self.T
        try:
            if not self.filament_var.get(): raise ValueError(T["no_filament"])
            grams=parse_number(self.grams.get(),T["total_filament_field"],T["invalid_number"])
            hours=parse_number(self.hours.get(),T["hours"],T["invalid_number"])
            minutes=parse_number(self.minutes.get(),T["minutes"],T["invalid_number"])
            if grams<0 or hours<0 or minutes<0: raise ValueError(T["negative_inputs"])
            if minutes>=60: raise ValueError(T["minutes_range"])
            match=next((r for r in self.db.filaments() if r[1]==self.filament_var.get()),None)
            if not match: raise ValueError(T["filament_not_found"])
            price_per_kg=float(match[2]); electricity=float(self.db.get_setting("electricity_price")); power_w=float(self.db.get_setting("printer_power_w"))
            total_hours=hours+minutes/60.0; material_cost=(grams/1000.0)*price_per_kg
            kwh=(power_w/1000.0)*total_hours; energy_cost=kwh*electricity; total=material_cost+energy_cost
            self.material_result.set(T["material_result"].format(cost=euro(material_cost), grams=grams, price=price_per_kg))
            self.kwh_result.set(T["kwh_result"].format(kwh=kwh,power=power_w,hours=total_hours))
            self.energy_result.set(T["energy_result"].format(cost=euro(energy_cost),electricity=electricity))
            self.total_result.set(euro(total))
            self.last_result={"filament":self.filament_var.get(),"grams":grams,"price_per_kg":price_per_kg,
                              "total_hours":total_hours,"power_w":power_w,"kwh":kwh,"electricity":electricity,
                              "material_cost":material_cost,"energy_cost":energy_cost,"total":total}
        except ValueError as e:
            messagebox.showerror(T["error"],str(e),parent=self)


if __name__ == "__main__":
    App().mainloop()
